"""重点产品故障 TOP3 统计（导出 Excel 的 sheet2）。

版式对齐《重点产品故障数量和类型.xlsx》：

    类别 | 型号 | 故障数量 | 占比 | TOP故障类型（单机型占比）

口径说明
--------
- **占比**：该型号故障数 ÷ 该 BU 全部故障数（0% 整数百分比，与参考表一致）。
  同一 BU 内只列头部机型，所以各机型占比之和会小于 100%，属正常。
- **TOP故障类型**：`1.故障类型--21%\\n2.故障类型--16%`，其中百分比是
  该故障类型 ÷ 该型号故障数（单机型内占比），至多 3 条，不足 3 条就少列。
- **A 列（类别）按 BU 纵向合并**，底色交替。
- 「未发现故障」不参与故障类型排名（但计入故障数量与分母）。
"""

from __future__ import annotations

import pandas as pd

KEY_SHEET_NAME = "重点产品故障TOP3"
KEY_COLUMNS = ["类别", "型号", "故障数量", "占比", "TOP故障类型（单机型占比）"]

DEFAULT_TOP_MODELS = 5      # 每个 BU 取故障数前 N 个型号
DEFAULT_TOP_FAULTS = 3      # 每个型号列前 N 个故障类型

# 不参与「故障类型」排名的值（计入数量，但不作为 TOP 故障）
EXCLUDE_FAULTS = ("未发现故障", "NFF", "无异常", "无异常（NFF）")

_EMPTY_TOKENS = ("", "nan", "nat", "none", "null", "n/a", "-")


# ==================== 小工具 ====================

def _s(value) -> str:
    """安全转字符串：NaN / None / 空 → ''。"""
    if value is None:
        return ""
    s = str(value).strip()
    return "" if s.lower() in _EMPTY_TOKENS else s


def _norm_model(name) -> str:
    """型号归一（复用 model_norm，兼容大小写/空格/下划线写法差异）。"""
    raw = _s(name)
    if not raw:
        return ""
    try:
        from model_norm import normalize_model
        return normalize_model(raw) or raw
    except Exception:
        return raw


def _bu_from_model(model: str) -> str:
    """按 config.PRODUCT_BU_MAP 的前缀表推断 BU（最长前缀优先）。

    先按原样匹配；失败再按「去掉空格」的紧凑形式匹配，
    以便 Air Style Pro → AirStyle、Aqua 10s → Aqua10 这类写法差异也能命中。
    """
    if not model:
        return ""
    try:
        from config import PRODUCT_BU_MAP
    except Exception:
        return ""

    low = model.lower()
    best, best_len = "", 0
    for prefix, bu in PRODUCT_BU_MAP.items():
        p = str(prefix).lower()
        if low.startswith(p) and len(p) > best_len:
            best, best_len = bu, len(p)
    if best:
        return best

    compact = low.replace(" ", "")
    for prefix, bu in PRODUCT_BU_MAP.items():
        p = str(prefix).lower().replace(" ", "")
        if compact.startswith(p) and len(p) > best_len:
            best, best_len = bu, len(p)
    return best


def _fmt_top_faults(pairs) -> str:
    """[(故障类型, 数量, 占比), ...] → '1.xx--21%\\n2.yy--16%'。"""
    lines = []
    for i, (name, _cnt, pct) in enumerate(pairs, 1):
        lines.append(f"{i}.{name}--{pct:.0%}")
    return "\n".join(lines)


# ==================== 主计算 ====================

def build_key_products(
    df: pd.DataFrame,
    top_models: int = DEFAULT_TOP_MODELS,
    top_faults: int = DEFAULT_TOP_FAULTS,
    bu_col: str = "产品BU",
    model_col: str = "型号",
    fault_col: str = "故障类型（一级）",
    exclude_faults=EXCLUDE_FAULTS,
    bu_filter=None,
) -> pd.DataFrame:
    """生成「重点产品故障 TOP3」表。

    参数
    ----
    df : 已处理工单 DataFrame（需含 型号 / 产品BU / 故障类型（一级））
    top_models : 每个 BU 保留故障数前 N 的型号
    top_faults : 每个型号保留前 N 个故障类型
    bu_filter : 可选，只统计指定 BU（list / tuple / set）

    返回
    ----
    DataFrame，列 = KEY_COLUMNS（无数据时返回空表，列齐全）
    """
    empty = pd.DataFrame(columns=KEY_COLUMNS)
    if df is None or len(df) == 0:
        return empty

    d = pd.DataFrame({
        "bu": df[bu_col].map(_s) if bu_col in df.columns else "",
        "model": df[model_col].map(_norm_model) if model_col in df.columns else "",
        "fault": df[fault_col].map(_s) if fault_col in df.columns else "",
    })

    # 型号为空的行无法归组，直接丢弃
    d = d[d["model"] != ""]
    if len(d) == 0:
        return empty

    # BU 缺失时按型号前缀推断
    blank = d["bu"] == ""
    if blank.any():
        d.loc[blank, "bu"] = d.loc[blank, "model"].map(_bu_from_model)
    d["bu"] = d["bu"].replace("", "未分类")

    if bu_filter:
        wanted = {_s(b) for b in bu_filter}
        d = d[d["bu"].isin(wanted)]
        if len(d) == 0:
            return empty

    excluded = {_s(x) for x in (exclude_faults or ())}
    top_models = max(1, int(top_models or 1))
    top_faults = max(1, int(top_faults or 1))

    # BU 排序：故障总量降序，同量按名称
    bu_totals = d.groupby("bu").size()
    bu_order = sorted(bu_totals.items(), key=lambda kv: (-kv[1], kv[0]))

    rows = []
    for bu, bu_total in bu_order:
        g = d[d["bu"] == bu]
        model_counts = g.groupby("model").size()
        # 数量降序，同量按型号名升序（保证输出稳定）
        model_order = sorted(model_counts.items(), key=lambda kv: (-kv[1], kv[0]))

        for model, cnt in model_order[:top_models]:
            gm = g[g["model"] == model]
            fc = gm[gm["fault"] != ""]["fault"].value_counts()
            fc = fc[[k for k in fc.index if k not in excluded]]
            pairs = [(name, int(n), n / cnt) for name, n in fc.head(top_faults).items()]

            rows.append({
                "类别": bu,
                "型号": model,
                "故障数量": int(cnt),
                "占比": cnt / bu_total,
                "TOP故障类型（单机型占比）": _fmt_top_faults(pairs),
            })

    return pd.DataFrame(rows, columns=KEY_COLUMNS)


def summarize_key_products(key_df: pd.DataFrame) -> dict:
    """给界面用的摘要：BU 数 / 型号数 / 总故障数。"""
    if key_df is None or len(key_df) == 0:
        return {"bus": 0, "models": 0, "total": 0}
    return {
        "bus": int(key_df["类别"].nunique()),
        "models": int(len(key_df)),
        "total": int(pd.to_numeric(key_df["故障数量"], errors="coerce").fillna(0).sum()),
    }


# ==================== 写 Excel（含样式）====================

_HEADER_FILL = "4E83FD"
_GROUP_FILLS = ("BACEFD", "E1EAFF")
_BORDER_COLOR = "DEE0E3"
_FONT_NAME = "等线"


def _style_key_sheet(ws, key_df: pd.DataFrame) -> None:
    """按参考文件样式美化 sheet2（表头蓝底、A 列按 BU 合并、E 列换行）。"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color=_BORDER_COLOR)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(name=_FONT_NAME, size=10, bold=True, color="000000")
    body_font = Font(name=_FONT_NAME, size=10)
    bold_font = Font(name=_FONT_NAME, size=10, bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=False)
    wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)

    n_cols = len(KEY_COLUMNS)
    n_rows = len(key_df)

    # 表头
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[1].height = 24

    # 数据行
    group_fill_idx = 0
    prev_bu = None
    group_start = 2
    groups = []          # [(start_row, end_row, fill_idx), ...]

    for i in range(n_rows):
        r = i + 2
        bu = str(key_df.iloc[i]["类别"])
        top_text = str(key_df.iloc[i]["TOP故障类型（单机型占比）"] or "")
        lines = max(1, top_text.count("\n") + 1)

        if bu != prev_bu:
            if prev_bu is not None:
                groups.append((group_start, r - 1, group_fill_idx))
                group_fill_idx += 1
            prev_bu, group_start = bu, r

        for c in range(1, n_cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = bold_font if c == 1 else body_font
            cell.alignment = wrap if c == n_cols else center
            cell.border = border
        # 占比列按 0% 显示（底层仍是小数）
        ws.cell(row=r, column=4).number_format = "0%"
        # 行高随 TOP 条数自适应
        ws.row_dimensions[r].height = {1: 20, 2: 30}.get(lines, 40)

    if prev_bu is not None:
        groups.append((group_start, n_rows + 1, group_fill_idx))

    # A 列按 BU 合并 + 交替底色
    for start, end, idx in groups:
        fill = PatternFill("solid", fgColor=_GROUP_FILLS[idx % len(_GROUP_FILLS)])
        for r in range(start, end + 1):
            ws.cell(row=r, column=1).fill = fill
        if end > start:
            ws.merge_cells(start_row=start, start_column=1,
                           end_row=end, end_column=1)

    # 列宽（参考表：型号 ~137px、数量 ~87px、TOP ~447px）
    widths = {1: 11, 2: 20, 3: 11, 4: 9, 5: 58}
    for c, w in widths.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    ws.freeze_panes = "A2"


def write_key_sheet_only(key_df: pd.DataFrame, target) -> None:
    """只写「重点产品故障TOP3」一个 sheet（带样式）。

    target 可以是路径或 BytesIO。样式失败不影响数据写入。
    """
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        key_df.to_excel(writer, index=False, sheet_name=KEY_SHEET_NAME)
        try:
            _style_key_sheet(writer.sheets[KEY_SHEET_NAME], key_df)
        except Exception:
            pass


def write_excel_with_key_sheet(
    main_df: pd.DataFrame,
    target,
    main_sheet: str = "已处理",
    top_models: int = DEFAULT_TOP_MODELS,
    top_faults: int = DEFAULT_TOP_FAULTS,
    key_df: pd.DataFrame | None = None,
    add_key_sheet: bool = True,
) -> pd.DataFrame:
    """写 Excel：sheet1 = 工单明细，sheet2 = 重点产品故障 TOP3。

    target 可以是路径，也可以是 BytesIO。返回本次写入的 key_df（可能为空表）。
    样式美化失败不影响数据写入（降级为无样式）。
    """
    if key_df is None and add_key_sheet:
        key_df = build_key_products(main_df, top_models=top_models,
                                    top_faults=top_faults)
    if key_df is None:
        key_df = pd.DataFrame(columns=KEY_COLUMNS)

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        main_df.to_excel(writer, index=False, sheet_name=main_sheet)
        if add_key_sheet and len(key_df) > 0:
            key_df.to_excel(writer, index=False, sheet_name=KEY_SHEET_NAME)
            try:
                _style_key_sheet(writer.sheets[KEY_SHEET_NAME], key_df)
            except Exception:
                pass   # 样式失败不阻断导出
    return key_df


if __name__ == "__main__":
    import argparse
    import io

    ap = argparse.ArgumentParser(description="从已处理工单生成重点产品故障 TOP3")
    ap.add_argument("xlsx", help="已处理工单 .xlsx / .csv")
    ap.add_argument("-o", "--out", default="重点产品故障TOP3.xlsx")
    ap.add_argument("--top-models", type=int, default=DEFAULT_TOP_MODELS)
    ap.add_argument("--top-faults", type=int, default=DEFAULT_TOP_FAULTS)
    args = ap.parse_args()

    src = args.xlsx
    raw = (pd.read_csv(src, dtype=str, encoding="utf-8-sig")
           if src.lower().endswith((".csv", ".tsv"))
           else pd.read_excel(src, dtype=str))
    kdf = build_key_products(raw, top_models=args.top_models,
                             top_faults=args.top_faults)
    write_excel_with_key_sheet(raw, args.out, main_sheet="已处理", key_df=kdf)
    print(f"已写出 {args.out}：{len(kdf)} 行")
    print(kdf.to_string(index=False))
